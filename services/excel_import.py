"""Parser do Excel de planilhas da Coach Pâm.

Lê o formato que ela já usa (cada aba = um time/fase) e devolve, por aba:
  - title, athletes (nomes do topo), weeks (qtd de semanas detectadas)
  - plan: { weeks: [ { week_number, days: [ { day, rest, lpo_exercises } ] } ] }
  - exercises: lista de nomes distintos encontrados
  - complexes: lista de complexos (nome com "+") com as partes

NÃO toca no banco. A resolução de exercise_id por nome e a validação contra a
biblioteca acontecem no router (que tem a sessão do DB).

Regras do layout (combinadas com a coach):
  - Cada semana é um bloco horizontal de 6 colunas:
        Exercício | RM | % | kg | s/r | Reps
    Os blocos começam onde a linha de cabeçalho tem "Exercício".
  - Dias ficam na 1ª coluna ("SEGUNDA-FEIRA", ... "QUINTA — DESCANSO").
  - Cabeçalhos decorativos (▸, ✦, ℹ️, emojis), PSE/PSR, AUXILIARES e TESTE
    são ignorados. A fonte de verdade é a LINHA DE DADOS (nome limpo + %).
  - "s/r" tipo "2x3" = 2 séries de 3 reps. Várias linhas = ondas (% diferentes).
  - "%" vem como decimal (0.70) e vira 70.
  - kg/RM/Reps(laranja)/PSE são ignorados no import (o app recalcula o kg
    pelo RM de cada atleta; Reps/PSE são dados de execução).
"""
from __future__ import annotations

import io
import re
from typing import List, Optional

import openpyxl


# ---- Mapeamento de dias -------------------------------------------------

_DAY_MAP = [
    (("segunda",), "monday"),
    (("terca", "terça"), "tuesday"),
    (("quarta",), "wednesday"),
    (("quinta",), "thursday"),
    (("sexta",), "friday"),
    (("sabado", "sábado"), "saturday"),
    (("domingo",), "sunday"),
]

# Prefixos/decorações que indicam linha que NÃO é dado de exercício.
_DECO_PREFIXES = ("▸", "✦", "ℹ", "🔵", "🟢", "💡", "👥", "🟡", "🟠", "⚠")


def _clean_text(v) -> str:
    return str(v).strip() if v is not None else ""


def _strip_accents_lower(s: str) -> str:
    s = (s or "").lower()
    repl = {"á": "a", "â": "a", "ã": "a", "à": "a", "é": "e", "ê": "e",
            "í": "i", "ó": "o", "ô": "o", "õ": "o", "ú": "u", "ç": "c"}
    for k, v in repl.items():
        s = s.replace(k, v)
    return s


def _detect_day(text: str) -> Optional[str]:
    """Se o texto da 1ª coluna marca um dia, retorna o slug do dia."""
    t = _strip_accents_lower(text)
    for needles, slug in _DAY_MAP:
        if any(n in t for n in needles):
            return slug
    return None


def _is_rest(text: str) -> bool:
    return "descanso" in _strip_accents_lower(text)


def _is_deco_or_meta(text: str) -> bool:
    """True para linhas decorativas/meta que devem ser ignoradas."""
    t = text.strip()
    if not t:
        return True
    if t.lower() == "x":
        return True
    for p in _DECO_PREFIXES:
        if t.startswith(p):
            return True
    tl = _strip_accents_lower(t)
    meta_markers = ("pse sessao", "psr", "auxiliares", "teste", "sem treino")
    if any(m in tl for m in meta_markers):
        return True
    return False


def parse_sr(sr) -> Optional[tuple]:
    """'2x3' -> (2, 3). Aceita também 'NxM' com espaços. None se não casar."""
    if sr is None:
        return None
    s = str(sr).strip().lower().replace(" ", "")
    m = re.match(r"^(\d+)x(\d+)$", s)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def pct_to_int(v) -> Optional[float]:
    """0.7 -> 70 ; 0.875 -> 87.5 ; 70 -> 70 (se já vier inteiro)."""
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f <= 1.5:  # decimal (0.70). Acima disso assume que já é percentual.
        f = f * 100
    r = round(f, 1)
    return int(r) if r == int(r) else r


def _strip_parens(name: str) -> str:
    """Tira observações entre parênteses, ex. '(2+1 = 1 rep)'."""
    return re.sub(r"\([^)]*\)", " ", name or "").strip()


def is_complex_name(name: str) -> bool:
    # "+" dentro de parênteses é esquema de reps, não complexo.
    return "+" in _strip_parens(name)


def split_complex(name: str) -> List[str]:
    parts = [p.strip() for p in re.split(r"\+", _strip_parens(name))]
    return [p for p in parts if p]


# ---- Detecção dos blocos de semana --------------------------------------

def _find_week_blocks(ws) -> List[int]:
    """Acha as colunas iniciais de cada bloco de semana.

    Procura na linha de cabeçalho (primeiras ~10 linhas) as células
    que contenham 'Exercício'. Cada uma marca o começo de um bloco.
    """
    for row in range(1, 11):
        cols = []
        for col in range(1, ws.max_column + 1):
            val = _clean_text(ws.cell(row=row, column=col).value)
            if _strip_accents_lower(val).startswith("exercicio"):
                cols.append(col)
        if len(cols) >= 1:
            return cols
    return [1]


def _find_header_row(ws) -> int:
    for row in range(1, 11):
        for col in range(1, ws.max_column + 1):
            val = _clean_text(ws.cell(row=row, column=col).value)
            if _strip_accents_lower(val).startswith("exercicio"):
                return row
    return 5


def _parse_athletes(ws) -> List[str]:
    """Nomes do time ficam numa linha com '👥 A · B · C'.

    Só consideramos a célula que tem o marcador 👥 (a do título também usa
    '·', então não dá pra cair nela).
    """
    for row in range(1, 8):
        for col in range(1, ws.max_column + 1):
            val = _clean_text(ws.cell(row=row, column=col).value)
            if "👥" in val:
                cleaned = val.replace("👥", "").strip()
                parts = re.split(r"·|,|;|/|\|", cleaned)
                names = [p.strip() for p in parts if p.strip()]
                if names:
                    return names
    return []


def _parse_title(ws) -> str:
    for row in range(1, 4):
        for col in range(1, ws.max_column + 1):
            val = _clean_text(ws.cell(row=row, column=col).value)
            if val and len(val) > 8 and "exercicio" not in _strip_accents_lower(val):
                # pega só a 1ª linha do texto do título
                return val.split("\n")[0].strip()
    return ws.title


def parse_sheet(ws) -> dict:
    """Parseia uma aba inteira e devolve a estrutura pronta para a prévia."""
    header_row = _find_header_row(ws)
    block_cols = _find_week_blocks(ws)
    first_data_row = header_row + 1
    last_row = ws.max_row

    # 1) Estrutura de dias a partir da 1ª coluna (col 1).
    #    Lista de (day_slug, rest, start_row, end_row).
    day_spans = []
    cur = None  # (slug, rest, start_row)
    for r in range(first_data_row, last_row + 1):
        a = _clean_text(ws.cell(row=r, column=1).value)
        slug = _detect_day(a)
        if slug:
            if cur:
                day_spans.append((cur[0], cur[1], cur[2], r - 1))
            cur = (slug, _is_rest(a), r + 1)
    if cur:
        day_spans.append((cur[0], cur[1], cur[2], last_row))

    weeks = []
    all_names: List[str] = []

    for wi, start_col in enumerate(block_cols):
        name_col = start_col
        pct_col = start_col + 2
        sr_col = start_col + 4

        days_out = []
        for slug, rest, r0, r1 in day_spans:
            lpo_exercises = []
            current_ex = None

            if not rest:
                for r in range(r0, r1 + 1):
                    raw_name = _clean_text(ws.cell(row=r, column=name_col).value)
                    pct = ws.cell(row=r, column=pct_col).value
                    sr = ws.cell(row=r, column=sr_col).value
                    pct_val = pct_to_int(pct)
                    sr_parsed = parse_sr(sr)

                    is_data_name = (
                        raw_name
                        and not _is_deco_or_meta(raw_name)
                        and _detect_day(raw_name) is None
                    )

                    if is_data_name:
                        # novo exercício (linha-base)
                        current_ex = {
                            "name": raw_name,
                            "is_complex": is_complex_name(raw_name),
                            "series": [],
                        }
                        lpo_exercises.append(current_ex)
                        all_names.append(raw_name)
                        _append_series(current_ex, pct_val, sr_parsed)
                    elif (pct_val is not None or sr_parsed is not None) and current_ex:
                        # onda adicional do exercício atual
                        _append_series(current_ex, pct_val, sr_parsed)

            # remove exercícios que ficaram sem nenhuma série nesta semana
            lpo_exercises = [e for e in lpo_exercises if e["series"]]
            days_out.append({
                "day": slug,
                "rest": rest,
                "coach_notes": "",
                "lpo_exercises": lpo_exercises,
                "accessories": [],
            })

        weeks.append({"week_number": wi + 1, "days": days_out})

    # nomes distintos preservando ordem
    seen = set()
    distinct = []
    for n in all_names:
        if n not in seen:
            seen.add(n)
            distinct.append(n)

    complexes = []
    for n in distinct:
        if is_complex_name(n):
            complexes.append({"name": n, "parts": split_complex(n)})

    return {
        "name": ws.title,
        "title": _parse_title(ws),
        "athletes": _parse_athletes(ws),
        "weeks": len(block_cols),
        "plan": {"weeks": weeks},
        "exercises": distinct,
        "complexes": complexes,
    }


def _append_series(ex: dict, pct_val, sr_parsed) -> None:
    """Expande s/r em séries. '2x3' -> 2 séries de {reps:3, percent}."""
    if sr_parsed:
        sets, reps = sr_parsed
    else:
        # sem s/r: 1 série; reps desconhecida -> 1 (a coach ajusta)
        sets, reps = 1, 1
    for _ in range(sets):
        ex["series"].append({
            "reps": reps,
            "percent": pct_val if pct_val is not None else 0,
        })


# =========================================================================
# Formato 2: "DIA N" + "Base do %"
#
#   DIA 1
#   #1  Back Squat na caixa        | SEMANA 1 | ... | SEMANA 4 | Observações
#   Exercício | Base do % | RM | %  | kg | Séries | % | kg | Séries | ...
#   Back Squat na caixa | Back Squat | | 0.75 | — | 3x5 | 0.8 | ...
#   (ou várias linhas "Série 1/2/3" = ondas)
#
# "Base do %" = movimento cujo RM é usado no cálculo do kg.
# =========================================================================

_DAY_RE = re.compile(r"^\s*dia\s*(\d+)\b", re.IGNORECASE)
_EX_HEADER_RE = re.compile(r"^\s*#\s*\d+\s*(.*)$")

# Qual dia da semana recebe cada "DIA N", conforme quantos dias o ciclo tem.
DEFAULT_DAY_SLOTS = {
    1: ["monday"],
    2: ["monday", "thursday"],
    3: ["monday", "wednesday", "friday"],
    4: ["monday", "tuesday", "thursday", "friday"],
    5: ["monday", "tuesday", "wednesday", "thursday", "friday"],
    6: ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday"],
    7: ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"],
}
_WEEK_ORDER = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]


def is_dia_format(ws) -> bool:
    has_dia = has_base = False
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), max_col=min(ws.max_column, 6)):
        for cell in row:
            t = _clean_text(cell.value)
            if not t:
                continue
            if cell.column == 1 and _DAY_RE.match(t):
                has_dia = True
            if _strip_accents_lower(t).startswith("base do"):
                has_base = True
        if has_dia and has_base:
            return True
    return False


def parse_sheet_dias(ws, day_slots: Optional[List[str]] = None) -> dict:
    """Parseia o formato 'DIA N' + 'Base do %'."""
    days: List[dict] = []          # [{"n": 1, "exercises": [...]}]
    cur_day = None
    cur_ex = None
    week_cols: List[int] = []      # coluna do % de cada semana
    notes_col: Optional[int] = None
    base_col = 2

    for r in range(1, ws.max_row + 1):
        a = _clean_text(ws.cell(row=r, column=1).value)

        m_day = _DAY_RE.match(a) if a else None
        if m_day:
            cur_day = {"n": int(m_day.group(1)), "exercises": []}
            days.append(cur_day)
            cur_ex = None
            continue

        m_ex = _EX_HEADER_RE.match(a) if a else None
        if m_ex and cur_day is not None:
            title = m_ex.group(1).strip()
            week_cols = []
            notes_col = None
            for c in range(2, ws.max_column + 1):
                v = _strip_accents_lower(_clean_text(ws.cell(row=r, column=c).value))
                if v.startswith("semana"):
                    week_cols.append(c)
                elif v.startswith("observa"):
                    notes_col = c
            cur_ex = {
                "name": title,
                "is_complex": is_complex_name(title),
                "rm_base": None,
                "weeks": [[] for _ in week_cols],
                "notes": [],
            }
            cur_day["exercises"].append(cur_ex)
            continue

        if cur_ex is None or not week_cols:
            continue

        # linha de cabeçalho da tabela
        if _strip_accents_lower(a).startswith("exercicio"):
            for c in range(1, ws.max_column + 1):
                if _strip_accents_lower(_clean_text(ws.cell(row=r, column=c).value)).startswith("base do"):
                    base_col = c
            continue

        base = _clean_text(ws.cell(row=r, column=base_col).value)
        if base and not cur_ex["rm_base"]:
            cur_ex["rm_base"] = base

        any_data = False
        for wi, pc in enumerate(week_cols):
            pct_val = pct_to_int(ws.cell(row=r, column=pc).value)
            sr = parse_sr(ws.cell(row=r, column=pc + 2).value)
            if pct_val is None and sr is None:
                continue
            any_data = True
            sets, reps = sr if sr else (1, 1)
            for _ in range(sets):
                cur_ex["weeks"][wi].append({
                    "reps": reps,
                    "percent": pct_val if pct_val is not None else 0,
                })

        if notes_col:
            note = _clean_text(ws.cell(row=r, column=notes_col).value)
            if note:
                cur_ex["notes"].append(note)

        if not any_data and not base and not a:
            # linha vazia encerra o bloco do exercício
            cur_ex = None

    # ---- monta o plano no formato do app ----
    n_weeks = max((len(ex["weeks"]) for d in days for ex in d["exercises"]), default=0)
    slots = day_slots or DEFAULT_DAY_SLOTS.get(len(days)) or _WEEK_ORDER[: len(days)]
    slot_by_n = {d["n"]: slots[i] for i, d in enumerate(days) if i < len(slots)}

    weeks_out = []
    for wi in range(n_weeks):
        days_out = []
        for slug in _WEEK_ORDER:
            day = next((d for d in days if slot_by_n.get(d["n"]) == slug), None)
            if day is None:
                days_out.append({"day": slug, "rest": True, "coach_notes": "",
                                 "lpo_exercises": [], "accessories": []})
                continue
            lpo = []
            notes = []
            for ex in day["exercises"]:
                series = ex["weeks"][wi] if wi < len(ex["weeks"]) else []
                if not series:
                    continue
                lpo.append({
                    "name": ex["name"],
                    "is_complex": ex["is_complex"],
                    "rm_base": ex["rm_base"],
                    "series": [dict(s) for s in series],
                })
                for n in ex["notes"]:
                    notes.append(f"{ex['name']}: {n}")
            days_out.append({"day": slug, "rest": False, "coach_notes": "\n".join(notes),
                             "lpo_exercises": lpo, "accessories": []})
        weeks_out.append({"week_number": wi + 1, "days": days_out})

    names = []
    for d in days:
        for ex in d["exercises"]:
            if ex["name"] not in names:
                names.append(ex["name"])
    bases = []
    for d in days:
        for ex in d["exercises"]:
            if ex["rm_base"] and ex["rm_base"] not in bases:
                bases.append(ex["rm_base"])

    title = ""
    for r in range(1, 4):
        v = _clean_text(ws.cell(row=r, column=1).value)
        if v:
            title = v if not title else f"{title} · {v}" if "ciclo" in _strip_accents_lower(v) else title
    return {
        "name": ws.title,
        "title": title or ws.title,
        "athletes": [],
        "weeks": n_weeks,
        "plan": {"weeks": weeks_out},
        "exercises": names,
        "bases": bases,
        "complexes": [{"name": n, "parts": split_complex(n)} for n in names if is_complex_name(n)],
        "format": "dias",
        "day_map": {f"DIA {d['n']}": slot_by_n.get(d["n"]) for d in days},
    }


def parse_workbook(file_bytes: bytes) -> List[dict]:
    """Parseia todas as abas (detectando o formato) e devolve as que têm treino."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    out = []
    for ws in wb.worksheets:
        parsed = parse_sheet_dias(ws) if is_dia_format(ws) else parse_sheet(ws)
        has_ex = any(
            d.get("lpo_exercises")
            for w in parsed["plan"]["weeks"] for d in w["days"]
        )
        if has_ex:
            out.append(parsed)
    return out
